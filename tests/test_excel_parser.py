from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from backend.indexing.document_loader import DocumentArtifactMetadata, DocumentLoader
from backend.indexing.excel_parser import parse_excel_rows
from backend.indexing.excel_provenance import excel_provenance, source_anchor


def workbook(tmp_path, rows, *, name="参数 !%中文"):
    book = Workbook()
    sheet = book.active
    sheet.title = name
    for row in rows:
        sheet.append(row)
    path = tmp_path / "table.xlsx"
    book.save(path)
    return path


def metadata(version="version-a"):
    return DocumentArtifactMetadata(
        "tenant-a", "kb-a", "doc-a", version, acl_tags=["reader"]
    )


@pytest.mark.parametrize(
    "value,expected",
    [
        (0, "0"),
        (-2, "-2"),
        (1.25, "1.25"),
        (True, "true"),
        (False, "false"),
        (date(2026, 1, 2), "2026-01-02"),
        (None, "（空）"),
    ],
)
def test_headers_typed_values_and_physical_rows(tmp_path, value, expected):
    path = workbook(tmp_path, [["型号", "数值"], ["NX-001", value]])
    row = parse_excel_rows(str(path))[0]
    assert f"型号：NX-001；数值：{expected}" in row["text"]
    assert row["row_start"] == row["row_end"] == 2


def test_empty_rows_duplicate_models_and_multisheet(tmp_path):
    path = workbook(tmp_path, [["型号", "功率W"], ["NX", 0], [None, None], ["NX", -1]])
    from openpyxl import load_workbook

    book = load_workbook(path)
    book.create_sheet("second").append(["型号", "功率W"])
    book["second"].append(["NX", 2])
    book.create_sheet("hidden").sheet_state = "hidden"
    book["hidden"].append(["型号", "功率W"])
    book["hidden"].append(["SECRET", 999])
    book.save(path)
    rows = parse_excel_rows(str(path))
    assert [row["row_start"] for row in rows] == [2, 4, 2]
    assert {row["sheet_name"] for row in rows} == {"参数 !%中文", "second"}
    assert "SECRET" not in str(rows)
    chunks = DocumentLoader().load_document(str(path), path.name, metadata=metadata())
    by_id = {chunk["chunk_id"]: chunk for chunk in chunks}
    for chunk in chunks:
        assert chunk["page_number"] == 0
        assert excel_provenance(chunk)
        if chunk["parent_chunk_id"]:
            assert by_id[chunk["parent_chunk_id"]]["sheet_name"] == chunk["sheet_name"]


@pytest.mark.parametrize(
    "rows",
    [
        [["型号", "值"], ["NX", "=1+2"]],
        [["型号", None], ["NX", 1]],
        [["值", "值"], [1, 2]],
        [["型号", "值"], ["NX", "x" * 10001]],
    ],
)
def test_unsupported_input_fails_closed(tmp_path, rows):
    with pytest.raises(ValueError):
        parse_excel_rows(str(workbook(tmp_path, rows)))


def test_long_row_fragments_repeat_identity_and_keep_row(tmp_path):
    path = workbook(tmp_path, [["型号", "描述"], ["NX", "x" * 2000]])
    rows = parse_excel_rows(str(path))
    assert len(rows) > 1
    assert all(
        "型号：NX；描述：" in row["text"] and len(row["text"]) <= 800 for row in rows
    )
    assert {row["row_start"] for row in rows} == {2}


def test_chunks_deterministic_and_version_scoped(tmp_path):
    path = workbook(tmp_path, [["型号", "功率W"], ["NX", 120]])
    loader = DocumentLoader()
    a = loader.load_document(str(path), path.name, metadata=metadata())
    assert a == loader.load_document(str(path), path.name, metadata=metadata())
    b = loader.load_document(str(path), path.name, metadata=metadata("version-b"))
    assert not {c["chunk_id"] for c in a} & {c["chunk_id"] for c in b}
    assert source_anchor("参数 !%中文", 2, 2) == a[-1]["source_anchor"]
    assert "%" in a[-1]["source_anchor"] and "参数" not in a[-1]["source_anchor"]


@pytest.mark.parametrize(
    "filename,expected",
    [
        ("product_specs.xlsx", 200),
        ("product_prices.xlsx", 200),
        ("inventory.xlsx", 200),
        ("repair_records.xlsx", 100),
    ],
)
def test_s1_real_workbooks_every_leaf_has_headers(filename, expected):
    path = (
        Path(__file__).resolve().parents[1]
        / "evals/rag/synthetic_enterprise/upload_corpus"
        / filename
    )
    chunks = DocumentLoader().load_document(str(path), filename, metadata=metadata())
    leaves = [chunk for chunk in chunks if chunk["chunk_level"] == 3]
    assert len(leaves) == expected
    assert len({chunk["chunk_id"] for chunk in chunks}) == len(chunks)
    assert all(
        "：" in chunk["text"]
        and chunk["row_start"] == chunk["row_end"]
        and excel_provenance(chunk)
        for chunk in leaves
    )


@pytest.mark.parametrize(
    "field,value",
    [
        ("row_start", True),
        ("row_end", "2"),
        ("source_anchor", "fake"),
        ("sheet_name", ""),
    ],
)
def test_provenance_type_validation(field, value):
    row = {
        "sheet_name": "s",
        "row_start": 2,
        "row_end": 2,
        "source_anchor": source_anchor("s", 2, 2),
    }
    with pytest.raises(ValueError):
        excel_provenance({**row, field: value})


def test_merged_headers_fail_explicitly(tmp_path):
    from openpyxl import load_workbook

    path = workbook(tmp_path, [["型号", "值"], ["NX", 1]])
    book = load_workbook(path)
    book.active.merge_cells("A1:B1")
    book.save(path)
    with pytest.raises(ValueError, match="merged"):
        parse_excel_rows(str(path))


@pytest.mark.parametrize(
    "limit",
    [
        "MAX_FILE_BYTES",
        "MAX_UNCOMPRESSED_BYTES",
        "MAX_ROWS",
        "MAX_COLUMNS",
        "MAX_CELLS",
    ],
)
def test_resource_limits_fail_closed(tmp_path, monkeypatch, limit):
    import backend.indexing.excel_parser as parser

    path = workbook(tmp_path, [["型号", "值"], ["NX", 1]])
    monkeypatch.setattr(parser, limit, 1)
    with pytest.raises(ValueError, match="limit"):
        parser.parse_excel_rows(str(path))


def test_legacy_xls_never_enters_xlsx_parser(tmp_path):
    with pytest.raises(ValueError, match="only accepts XLSX"):
        parse_excel_rows(str(tmp_path / "old.xls"))


def test_underreported_dimensions_do_not_silently_lose_rows(tmp_path):
    from zipfile import ZipFile

    original = workbook(tmp_path, [["型号", "值"], ["NX", 1]])
    malformed = tmp_path / "malformed.xlsx"
    with ZipFile(original) as source, ZipFile(malformed, "w") as target:
        for member in source.infolist():
            content = source.read(member)
            if member.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="A1:B2"', b'ref="A1:A1"')
            target.writestr(member, content)
    with pytest.raises(ValueError, match="dimensions omit"):
        parse_excel_rows(str(malformed))
