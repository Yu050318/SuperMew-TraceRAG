from collections import Counter
from io import BytesIO
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from backend.evaluation.rag import RagEvalDataset, load_rag_eval_dataset
from scripts.generate_synthetic_enterprise import (
    SEED,
    TABLES,
    build_assets,
    generate,
    make_cases,
    make_facts,
    make_records,
    sha256,
    validate_assets,
)


@pytest.fixture(scope="module")
def assets():
    return build_assets()


def test_seed_reproduces_all_bytes(assets):
    assert assets == build_assets(SEED)
    assert make_records(SEED) != make_records(SEED + 1)


def test_product_keys_and_repair_foreign_keys():
    records = make_records()
    assert [p["id"] for p in records["products"]] == [
        f"NX-{i:03d}" for i in range(1, 201)
    ]
    assert len(records["repairs"]) == 100
    assert all(
        r["product_id"] in {p["id"] for p in records["products"]}
        for r in records["repairs"]
    )


@pytest.mark.parametrize("file", TABLES)
def test_xlsx_physical_rows_and_foreign_keys(assets, file):
    book = load_workbook(BytesIO(assets[f"corpus/{file}"]), data_only=True)
    sheet, headers = TABLES[file]
    rows = list(book[sheet].values)
    assert list(rows[0]) == headers
    assert book[sheet].max_row == (101 if file == "repair_records.xlsx" else 201)
    column = 1 if file == "repair_records.xlsx" else 0
    assert [r[column] for r in rows[1:]] == [f"NX-{i:03d}" for i in range(1, len(rows))]
    assert all(cell.data_type != "f" for row in book[sheet] for cell in row)
    book.close()


@pytest.mark.parametrize(
    "category,count,dev",
    [
        ("simple_fact", 20, 5),
        ("table_query", 20, 5),
        ("cross_document", 20, 5),
        ("no_answer", 10, 3),
        ("version_policy", 10, 2),
    ],
)
def test_category_and_stratified_splits(assets, category, count, dev):
    data = json.loads(assets["dataset_v1.json"])
    cases = [c for c in data["cases"] if category in c["tags"]]
    assert len(cases) == count
    assert sum("dev" in c["tags"] for c in cases) == dev


@pytest.mark.parametrize(
    "name,count",
    [
        ("dataset_v1.json", 80),
        ("dataset_dev_v1.json", 20),
        ("dataset_frozen_v1.json", 60),
    ],
)
def test_native_dataset_contract(assets, name, count):
    data = RagEvalDataset.model_validate_json(assets[name])
    assert len(data.cases) == count
    assert all(not c.gold_chunks for c in data.cases)
    assert all(d.document_id is None for c in data.cases for d in c.gold_documents)


def test_all_answers_derive_from_source_and_sources_exist(assets):
    records = json.loads(assets["source_records.json"])
    data, truth = make_cases(records, make_facts(records))
    assert data == json.loads(assets["dataset_v1.json"])
    assert truth == [
        json.loads(line) for line in assets["ground_truth.jsonl"].splitlines()
    ]
    validate_assets(
        records,
        data,
        truth,
        {
            name.removeprefix("corpus/"): value
            for name, value in assets.items()
            if name.startswith("corpus/")
        },
    )


def test_all_references_and_ids_are_unique_and_bound(assets):
    data = json.loads(assets["dataset_v1.json"])
    assert len({c["id"] for c in data["cases"]}) == 80
    manifest = json.loads(assets["manifest.json"])
    names = {Path(n).name for n in manifest["ingestion_allowlist"]}
    assert len(names) == 12
    for case in data["cases"]:
        references = [d["canonical_name"] for d in case["gold_documents"]]
        assert len(references) == len(set(references))
        assert set(references) <= names


def test_no_answer_models_not_in_corpus(assets):
    data = json.loads(assets["dataset_v1.json"])
    products = {p["id"] for p in json.loads(assets["source_records.json"])["products"]}
    text = " ".join(
        v.decode()
        for n, v in assets.items()
        if n.startswith("upload_corpus/") and n.endswith(".html")
    )
    for case in data["cases"]:
        if "no_answer" in case["tags"]:
            model = case["question"].split()[0]
            assert model not in products and model not in text
            assert not case["gold_documents"] and not case["required_claims"]
            assert case["expected"]["acceptable_abstention"]


def test_cross_document_claims_need_both_sources(assets):
    truth = [json.loads(line) for line in assets["ground_truth.jsonl"].splitlines()]
    for item in truth:
        if item["category"] == "cross_document":
            assert len({f["file"] for f in item["evidence"]}) == 2
            assert len(item["fact_ids"]) == 2
            assert all(
                f["statement"] in item["reference_answer"] for f in item["evidence"]
            )


def test_conflicts_have_explicit_distinct_version_rules(assets):
    records = json.loads(assets["source_records.json"])
    assert records["policy"]["old_months"] != records["policy"]["current_months"]
    products = {p["id"]: p for p in records["products"]}
    assert all(
        v["old_power_w"] != products[v["product_id"]]["power_w"]
        for v in records["versions"]
    )
    data = json.loads(assets["dataset_v1.json"])
    assert sum(bool(c["conflicts"]) for c in data["cases"]) == 10


def test_ingestion_allowlist_excludes_all_labels(assets):
    manifest = json.loads(assets["manifest.json"])
    assert len(manifest["ingestion_allowlist"]) == 12
    for name in manifest["ingestion_allowlist"]:
        assert name.startswith("upload_corpus/")
        assert Path(name).suffix in {".xlsx", ".html"}
        assert name in assets
        if name.endswith(".html"):
            assert b"reference_answer" not in assets[name]
            assert b"case_id" not in assets[name]
    assert all(
        not name.startswith("upload_corpus/")
        for name in ("ground_truth.jsonl", "source_records.json", "dataset_v1.json")
    )


def test_manifest_hashes_match_every_generated_asset(assets):
    manifest = json.loads(assets["manifest.json"])
    assert set(manifest["files"]) == set(assets) - {"manifest.json"}
    assert all(
        sha256(assets[name]) == digest for name, digest in manifest["files"].items()
    )


def test_dev_frozen_disjoint_and_cover_dataset(assets):
    splits = json.loads(assets["splits.json"])
    assert not set(splits["dev"]) & set(splits["frozen"])
    assert len(set(splits["dev"]) | set(splits["frozen"])) == 80
    assert Counter(
        c["split"] for c in map(json.loads, assets["ground_truth.jsonl"].splitlines())
    ) == {"dev": 20, "frozen": 60}


def test_regeneration_no_changes_and_loads_native(tmp_path):
    first = generate(tmp_path)
    mtimes = {p: p.stat().st_mtime_ns for p in tmp_path.rglob("*") if p.is_file()}
    assert generate(tmp_path) == first
    assert generate(tmp_path, check=True) == first
    assert all(p.stat().st_mtime_ns == stamp for p, stamp in mtimes.items())
    assert len(load_rag_eval_dataset(tmp_path / "dataset_v1.json").cases) == 80


@pytest.mark.parametrize(
    "target", ["source_records.json", "corpus/product_manual.md", "dataset_v1.json"]
)
def test_regeneration_refuses_hand_edits_before_writing(tmp_path, target):
    generate(tmp_path)
    (tmp_path / target).write_text("hand-authored data", encoding="utf-8")
    before = {p: p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    with pytest.raises(ValueError, match="refusing to overwrite"):
        generate(tmp_path)
    assert before == {p: p.read_bytes() for p in before}


def test_existing_manual_file_in_fresh_output_is_protected(tmp_path):
    (tmp_path / "source_records.json").write_text("manual", encoding="utf-8")
    with pytest.raises(ValueError, match="refusing to overwrite"):
        generate(tmp_path)
    assert not (tmp_path / "corpus").exists()


def test_extra_file_in_upload_corpus_rejected(tmp_path):
    generate(tmp_path)
    (tmp_path / "upload_corpus/answers.json").write_text("answers", encoding="utf-8")
    with pytest.raises(AssertionError, match="unexpected file"):
        generate(tmp_path, check=True)


def test_tampered_ground_truth_rejected(assets):
    records = json.loads(assets["source_records.json"])
    data, truth = make_cases(records, make_facts(records))
    truth[0]["reference_answer"] = "invented answer"
    with pytest.raises(AssertionError, match="derive from records"):
        validate_assets(records, data, truth, {})


def test_manual_review_has_ten_pending_samples(assets):
    assert assets["manual_review.md"].decode().count("- [ ]") == 10


@pytest.mark.parametrize(
    "name",
    [
        "product_manual",
        "aftersales_policy",
        "warranty_rules",
        "product_faq",
        "technical_specs",
        "release_notes",
        "troubleshooting",
        "selection_guide",
    ],
)
def test_upload_html_preserves_facts_through_production_parser(assets, tmp_path, name):
    pytest.importorskip("bs4")
    from backend.indexing.html_processor import parse_html_file_to_sections

    file = tmp_path / f"{name}.html"
    file.write_bytes(assets[f"upload_corpus/{name}.html"])
    sections = parse_html_file_to_sections(file)
    rendered = "\n".join(section["text"] for section in sections)
    facts = make_facts(json.loads(assets["source_records.json"]))
    for fact in facts:
        if fact["file"] == f"{name}.md":
            assert fact["statement"] in rendered
    assert sections and "虚构测试数据" in rendered
