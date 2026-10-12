"""XLSX v1: explicit first-row headers, bounded cells, no formula execution."""

from datetime import date, datetime, time
from pathlib import Path
import re
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 100
MAX_CELLS = 200000
MAX_CELL_CHARACTERS = 10000
MAX_SHEETS = 32


def _display(value: object) -> str:
    if value is None:
        return "（空）"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def parse_excel_rows(path: str, *, fragment_size: int = 800) -> list[dict]:
    from backend.indexing.document_loader import sanitize_text

    if Path(path).suffix.lower() != ".xlsx":
        # Stored upload paths may have a suffix but legacy XLS is never accepted here.
        raise ValueError("Excel parser only accepts XLSX")
    if not 600 <= fragment_size <= 4000:
        raise ValueError("invalid Excel fragment budget")
    if Path(path).stat().st_size > MAX_FILE_BYTES:
        raise ValueError("XLSX file size exceeds limit")
    with ZipFile(path) as archive:
        members = archive.infolist()
        if len({item.filename for item in members}) != len(members):
            raise ValueError("duplicate XLSX archive entries")
        if any("vbaProject" in item.filename for item in members):
            raise ValueError("XLSX macros are unsupported")
        if (
            len(members) > 2048
            or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES
        ):
            raise ValueError("XLSX expanded size exceeds limit")
        for member in members:
            if member.filename.startswith(
                "xl/worksheets/"
            ) and member.filename.endswith(".xml"):
                declared_column = declared_row = 0
                actual_column = actual_row = 0
                with archive.open(member) as stream:
                    for _, node in ElementTree.iterparse(stream, events=("end",)):
                        if node.tag.endswith("}dimension"):
                            endpoint = node.get("ref", "").split(":")[-1]
                            match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", endpoint)
                            if not match:
                                raise ValueError("invalid XLSX worksheet dimensions")
                            for letter in match[1]:
                                declared_column = (
                                    declared_column * 26 + ord(letter) - ord("A") + 1
                                )
                            declared_row = int(match[2])
                        if node.tag.endswith("}mergeCell"):
                            raise ValueError("merged cells are unsupported in XLSX v1")
                        if (
                            node.tag.endswith("}row")
                            and int(node.get("r", "0")) > MAX_ROWS
                        ):
                            raise ValueError("XLSX physical row exceeds limit")
                        if node.tag.endswith("}c"):
                            coordinate = node.get("r", "")
                            match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", coordinate)
                            if not match:
                                raise ValueError("invalid XLSX cell coordinate")
                            column = 0
                            for letter in match[1]:
                                column = column * 26 + ord(letter) - ord("A") + 1
                            if column > MAX_COLUMNS or int(match[2]) > MAX_ROWS:
                                raise ValueError("XLSX physical cell exceeds limit")
                            actual_column = max(actual_column, column)
                            actual_row = max(actual_row, int(match[2]))
                        node.clear()
                if actual_column > declared_column or actual_row > declared_row:
                    raise ValueError("XLSX dimensions omit physical cells")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    result: list[dict] = []
    cell_count = 0
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ValueError("XLSX sheet count exceeds limit")
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            if sheet.max_row is None or sheet.max_column is None:
                raise ValueError("XLSX requires bounded worksheet dimensions")
            if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
                raise ValueError("XLSX row/column count exceeds limit")
            cell_count += sheet.max_row * sheet.max_column
            if cell_count > MAX_CELLS:
                raise ValueError("XLSX cell count exceeds limit")
            headers: list[str] | None = None
            for number, cells in enumerate(sheet.iter_rows(), 1):
                values: list[str] = []
                for cell in cells:
                    if cell.data_type in {"f", "e"}:
                        raise ValueError(
                            "XLSX formulas/cell errors are unsupported; export computed values"
                        )
                    raw_text = _display(cell.value)
                    if len(raw_text) > MAX_CELL_CHARACTERS:
                        raise ValueError("XLSX cell text exceeds limit")
                    text = sanitize_text(raw_text).strip()
                    if len(text) > MAX_CELL_CHARACTERS:
                        raise ValueError("XLSX cell text exceeds limit")
                    values.append(text)
                if number == 1:
                    if all(cell.value is None for cell in cells) and sheet.max_row == 1:
                        break
                    if any(cell.value is None for cell in cells) or len(
                        set(values)
                    ) != len(values):
                        raise ValueError(
                            "XLSX requires unique non-empty first-row headers"
                        )
                    if any(not value or len(value) > 128 for value in values):
                        raise ValueError("XLSX header exceeds limit")
                    headers = values
                    continue
                if all(cell.value is None for cell in cells):
                    continue
                assert headers is not None
                # First-column identity is repeated on long-row fragments, not inferred as unique.
                identity = f"{headers[0]}：{values[0]}"
                if len(identity) > fragment_size // 2:
                    raise ValueError(
                        "XLSX first-column identity exceeds fragment budget"
                    )
                prefix = f"Sheet：{sheet.title}；{identity}"
                fragments: list[str] = []
                current = prefix
                for header, value in zip(headers[1:], values[1:]):
                    budget = fragment_size - len(prefix) - len(header) - 8
                    for offset in range(0, max(1, len(value)), budget):
                        field = f"{header}：{value[offset : offset + budget]}"
                        if len(current) + len(field) + 1 > fragment_size:
                            fragments.append(current)
                            current = prefix
                        current += f"；{field}"
                fragments.append(current)
                result.extend(
                    {
                        "sheet_name": sheet.title,
                        "row_start": number,
                        "row_end": number,
                        "text": text,
                    }
                    for text in fragments
                )
    finally:
        workbook.close()
    if not result:
        raise ValueError("XLSX has no visible data rows")
    return result
