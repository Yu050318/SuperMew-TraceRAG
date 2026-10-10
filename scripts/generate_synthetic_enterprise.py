"""Deterministic synthetic assets; only upload_corpus/ is an ingestion allowlist."""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime
import hashlib
from html import escape
from io import BytesIO
import json
from pathlib import Path
import random
import re
import sys
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backend.evaluation.rag import RagEvalDataset, dataset_fingerprint  # noqa: E402

VERSION = "nebula-v1"
SEED = 20261010
NOTICE = "SYNTHETIC DATA / 虚构测试数据 — 星云智能科技有限公司 (Nebula Industrial)"
DEFAULT_OUTPUT = ROOT / "evals/rag/synthetic_enterprise"
TABLES = {
    "product_specs.xlsx": (
        "产品参数",
        [
            "型号",
            "产品类别",
            "额定功率 (W)",
            "重量 (kg)",
            "协议",
            "硬件版本",
            "生效日期",
        ],
    ),
    "product_prices.xlsx": ("价格", ["型号", "单价 (CNY)", "币种", "生效日期"]),
    "inventory.xlsx": ("库存", ["型号", "仓库", "可用库存 (台)", "更新时间"]),
    "repair_records.xlsx": (
        "维修",
        ["维修编号", "型号", "故障代码", "处理结论", "维修日期"],
    ),
}
MARKDOWN = (
    "product_manual.md",
    "aftersales_policy.md",
    "warranty_rules.md",
    "product_faq.md",
    "technical_specs.md",
    "release_notes.md",
    "troubleshooting.md",
    "selection_guide.md",
)


def json_bytes(value: object) -> bytes:
    return (
        json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n"
    ).encode("utf-8")


def sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def make_records(seed: int = SEED) -> dict:
    rng = random.Random(seed)
    products = []
    for i in range(1, 201):
        products.append(
            {
                "id": f"NX-{i:03d}",
                "category": ("工业传感器", "工业控制器", "智能网关")[(i - 1) % 3],
                "power_w": rng.randrange(10, 201, 5),
                "weight_kg": rng.randrange(5, 81) / 10,
                "protocol": ("Modbus RTU", "EtherCAT", "MQTT")[(i - 1) % 3],
                "version": "2.0",
                "effective_date": "2026-01-01",
                "price_cny": rng.randrange(300, 6001, 50),
                "currency": "CNY",
                "stock": rng.randrange(0, 301),
                "warehouse": "上海测试仓",
                "stock_date": "2026-10-01",
                "voltage": "24 V DC",
                "installation": "断电后接线并固定在 DIN 导轨上",
                "maintenance": "每 6 个月检查接线端子",
                "environment": "室内干燥环境，禁止浸水",
                "recovery": "断电 30 秒后重新上电",
                "selection": "按协议与额定功率匹配，不按外观替代",
            }
        )
    repairs = [
        {
            "id": f"REP-{i:04d}",
            "product_id": products[i - 1]["id"],
            "code": "E01",
            "resolution": "重新紧固端子后恢复",
            "date": "2026-09-01",
        }
        for i in range(1, 101)
    ]
    return {
        "generator_version": VERSION,
        "seed": seed,
        "company": NOTICE,
        "products": products,
        "repairs": repairs,
        "policy": {
            "old_months": 12,
            "current_months": 24,
            "old_date": "2025-01-01",
            "current_date": "2026-01-01",
            "rule": "按购买日期适用，2026-01-01 起购买适用新政策",
        },
        "versions": [
            {
                "product_id": p["id"],
                "old_power_w": p["power_w"] + 10,
                "old_version": "1.0",
                "old_date": "2025-01-01",
            }
            for p in products[:5]
        ],
    }


def make_facts(records: dict) -> list[dict]:
    """Derive renderable facts and physical anchors from the sole business source."""
    facts = []

    def add(
        key: str,
        text: str,
        file: str,
        *,
        sheet: str | None = None,
        row: int | None = None,
        section: str | None = None,
    ) -> None:
        facts.append(
            {
                "id": key,
                "statement": text,
                "file": file,
                "sheet": sheet,
                "row_start": row,
                "row_end": row,
                "section": section,
                "source_anchor": f"{file}::{sheet or section}::{row or 'section'}",
            }
        )

    md_fields = {
        "manual": ("product_manual.md", "installation", "安装要求"),
        "faq": ("product_faq.md", "environment", "使用环境"),
        "technical": ("technical_specs.md", "voltage", "供电要求"),
        "trouble": ("troubleshooting.md", "recovery", "复位方法"),
        "selection": ("selection_guide.md", "selection", "选型规则"),
    }
    for row, p in enumerate(records["products"], 2):
        pid = p["id"]
        for field, label, unit in (
            ("power_w", "额定功率", "W"),
            ("weight_kg", "重量", "kg"),
            ("protocol", "通信协议", ""),
        ):
            add(
                f"{pid}.{field}",
                f"{pid} 的{label}为 {p[field]}{(' ' + unit) if unit else ''}。",
                "product_specs.xlsx",
                sheet="产品参数",
                row=row,
            )
        add(
            f"{pid}.price",
            f"{pid} 的单价为 {p['price_cny']} CNY。",
            "product_prices.xlsx",
            sheet="价格",
            row=row,
        )
        add(
            f"{pid}.stock",
            f"{pid} 在上海测试仓的可用库存为 {p['stock']} 台（2026-10-01）。",
            "inventory.xlsx",
            sheet="库存",
            row=row,
        )
        for kind, (file, field, label) in md_fields.items():
            add(f"{pid}.{kind}", f"{pid} 的{label}：{p[field]}。", file, section=pid)
        add(
            f"{pid}.maintenance",
            f"{pid} 的维护要求：{p['maintenance']}。",
            "warranty_rules.md",
            section=pid,
        )
    for row, r in enumerate(records["repairs"], 2):
        add(
            r["id"],
            f"维修 {r['id']} 对应 {r['product_id']}，故障 {r['code']}，结论：{r['resolution']}。",
            "repair_records.xlsx",
            sheet="维修",
            row=row,
        )
    policy = records["policy"]
    add(
        "policy.old",
        f"2025-01-01 至 2025-12-31 购买的设备保修 {policy['old_months']} 个月。",
        "aftersales_policy.md",
        section="2025政策",
    )
    add(
        "policy.current",
        f"2026-01-01 起购买的设备保修 {policy['current_months']} 个月。",
        "aftersales_policy.md",
        section="2026政策",
    )
    add("policy.rule", policy["rule"] + "。", "warranty_rules.md", section="适用规则")
    for v in records["versions"]:
        add(
            f"{v['product_id']}.old_power",
            f"{v['product_id']} 硬件 1.0 的额定功率为 {v['old_power_w']} W（2025-01-01）；2.0 自 2026-01-01 生效，以产品参数表的 2.0 数值为准。",
            "release_notes.md",
            section=v["product_id"],
        )
    return facts


def upload_name(file: str) -> str:
    return str(Path(file).with_suffix(".html")) if file.endswith(".md") else file


def make_cases(records: dict, facts: list[dict]) -> tuple[dict, list[dict]]:
    lookup = {f["id"]: f for f in facts}
    cases, truth = [], []
    counts = {
        "simple_fact": 20,
        "table_query": 20,
        "cross_document": 20,
        "no_answer": 10,
        "version_policy": 10,
    }
    dev_counts = {
        "simple_fact": 5,
        "table_query": 5,
        "cross_document": 5,
        "no_answer": 3,
        "version_policy": 2,
    }

    def add(
        category: str,
        number: int,
        question: str,
        keys: list[str],
        *,
        answer: str | None = None,
        conflicts: list[str] | None = None,
    ) -> None:
        case_id = f"{category}_{number:03d}"
        selected = [lookup[key] for key in keys]
        claims = [f["statement"] for f in selected]
        split = "dev" if number <= dev_counts[category] else "frozen"
        abstain = category == "no_answer"
        reference = answer or " ".join(claims)
        case = {
            "id": case_id,
            "tags": [category, split, "synthetic"],
            "question": question,
            "expected": {
                "outcome": "NO_KNOWLEDGE" if abstain else "ANSWERABLE",
                "acceptable_abstention": abstain,
            },
            "gold_documents": [
                {"canonical_name": name}
                for name in sorted({upload_name(f["file"]) for f in selected})
            ],
            "gold_chunks": [],
            "reference_answer": reference,
            "required_claims": claims,
            "conflicts": conflicts or [],
        }
        cases.append(case)
        truth.append(
            {
                "case_id": case_id,
                "category": category,
                "split": split,
                "fact_ids": keys,
                "reference_answer": reference,
                "evidence": selected,
                "conflict_expectation": conflicts or [],
                "unknown_product_id": f"NX-{900 + number:03d}" if abstain else None,
            }
        )

    for i in range(1, counts["simple_fact"] + 1):
        pid = records["products"][i + 39]["id"]
        kind, label = (
            ("manual", "安装要求"),
            ("faq", "使用环境"),
            ("technical", "供电要求"),
            ("trouble", "复位方法"),
        )[((i - 1) % 4)]
        add("simple_fact", i, f"{pid} 的{label}是什么？", [f"{pid}.{kind}"])
    for i in range(1, 21):
        p = records["products"][i - 1]
        pid = p["id"]
        if i <= 10:
            field, label = (
                ("power_w", "额定功率"),
                ("price", "单价"),
                ("stock", "库存"),
            )[(i - 1) % 3]
            add("table_query", i, f"{pid} 的{label}是多少？", [f"{pid}.{field}"])
        else:
            other = records["products"][i + 19]
            relation = (
                "大于"
                if p["power_w"] > other["power_w"]
                else "小于"
                if p["power_w"] < other["power_w"]
                else "等于"
            )
            keys = [f"{pid}.power_w", f"{other['id']}.power_w"]
            add(
                "table_query",
                i,
                f"{pid} 和 {other['id']} 的额定功率各是多少，谁更大？",
                keys,
                answer=" ".join(lookup[k]["statement"] for k in keys)
                + f" {pid} 的功率{relation} {other['id']}。",
            )
    for i in range(1, 21):
        pid = records["products"][i + 79]["id"]
        keys = (
            [f"{pid}.power_w", "policy.current"]
            if i % 2
            else [f"{pid}.protocol", f"{pid}.manual"]
        )
        question = (
            f"2026-02-01 购买的 {pid} 额定功率和保修期分别是多少？"
            if i % 2
            else f"{pid} 的通信协议和安装要求分别是什么？"
        )
        add("cross_document", i, question, keys)
    for i in range(1, 11):
        add(
            "no_answer",
            i,
            f"NX-{900 + i:03d} 的额定功率是多少？",
            [],
            answer="知识库未提供该型号信息，无法确认额定功率。",
        )
    for i in range(1, 11):
        if i <= 5:
            pid = records["products"][i - 1]["id"]
            add(
                "version_policy",
                i,
                f"{pid} 的公告与参数表功率不同；硬件 2.0 应使用哪个数值，如何解释差异？",
                [f"{pid}.power_w", f"{pid}.old_power"],
                conflicts=[
                    "公告的 1.0 功率与参数表 2.0 不同，应披露版本并选择 2.0 参数。"
                ],
            )
        else:
            add(
                "version_policy",
                i,
                f"售后说明同时写了 12 和 24 个月；2026-{i:02d}-01 购买适用哪个保修期，旧政策为何不同？",
                ["policy.old", "policy.current", "policy.rule"],
                conflicts=[
                    "披露 2025/2026 保修期差异，按购买日期适用 24 个月，不能把旧政策当现行规则。"
                ],
            )
    dataset = RagEvalDataset.model_validate(
        {"schema_version": 1, "name": "nebula_enterprise_v1", "cases": cases}
    )
    return dataset.model_dump(mode="json", exclude_none=True), truth


def table_rows(records: dict) -> dict[str, list[list]]:
    products = records["products"]
    return {
        "product_specs.xlsx": [
            [
                p[k]
                for k in (
                    "id",
                    "category",
                    "power_w",
                    "weight_kg",
                    "protocol",
                    "version",
                    "effective_date",
                )
            ]
            for p in products
        ],
        "product_prices.xlsx": [
            [p[k] for k in ("id", "price_cny", "currency", "effective_date")]
            for p in products
        ],
        "inventory.xlsx": [
            [p[k] for k in ("id", "warehouse", "stock", "stock_date")] for p in products
        ],
        "repair_records.xlsx": [
            [r[k] for k in ("id", "product_id", "code", "resolution", "date")]
            for r in records["repairs"]
        ],
    }


def xlsx_bytes(sheet: str, headers: list[str], rows: list[list]) -> bytes:
    book = Workbook()
    book.properties.creator = NOTICE
    book.properties.created = datetime(2026, 10, 10)
    book.properties.modified = datetime(2026, 10, 10)
    page = book.active
    page.title = sheet
    page.append(headers)
    for row in rows:
        page.append(row)
    raw = BytesIO()
    book.save(raw)
    # openpyxl sets modified and ZIP timestamps at save; normalize both for byte identity.
    output = BytesIO()
    with (
        ZipFile(raw) as source,
        ZipFile(output, "w", compression=ZIP_DEFLATED) as target,
    ):
        for name in sorted(source.namelist()):
            value = source.read(name)
            if name == "docProps/core.xml":
                value = re.sub(
                    rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                    rb"\g<1>2026-10-10T00:00:00Z\2",
                    value,
                )
            info = ZipInfo(name, date_time=(2026, 10, 10, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, value)
    return output.getvalue()


def validate_assets(
    records: dict, dataset: dict, truth: list[dict], corpus: dict[str, bytes]
) -> None:
    """Check foreign keys, native contract, exact cells, source support and split."""
    native = RagEvalDataset.model_validate(dataset)
    ids = [p["id"] for p in records["products"]]
    assert len(ids) == len(set(ids)) == 200
    assert all(r["product_id"] in ids for r in records["repairs"])
    facts = make_facts(records)
    expected_dataset, expected_truth = make_cases(records, facts)
    assert dataset == expected_dataset and truth == expected_truth, (
        "answers must derive from records"
    )
    assert len(native.cases) == 80
    assert Counter(t["split"] for t in truth) == {"dev": 20, "frozen": 60}
    assert Counter(t["category"] for t in truth) == {
        "simple_fact": 20,
        "table_query": 20,
        "cross_document": 20,
        "no_answer": 10,
        "version_policy": 10,
    }
    assert len({f["id"] for f in facts}) == len(facts)
    assert len(corpus) == 12 and set(corpus) == set(TABLES) | set(MARKDOWN)
    for file, rows in table_rows(records).items():
        sheet, headers = TABLES[file]
        book = load_workbook(BytesIO(corpus[file]), read_only=True, data_only=True)
        try:
            assert list(book[sheet].values) == [
                tuple(headers),
                *(tuple(row) for row in rows),
            ]
        finally:
            book.close()
    for fact in facts:
        assert fact["file"] in corpus
        if fact["file"].endswith(".md"):
            assert fact["statement"] in corpus[fact["file"]].decode("utf-8")
            assert f"## {fact['section']}\n" in corpus[fact["file"]].decode("utf-8")
        else:
            assert 2 <= fact["row_start"] <= len(table_rows(records)[fact["file"]]) + 1
    corpus_text = " ".join(p["id"] for p in records["products"]) + " ".join(
        v.decode("utf-8") for k, v in corpus.items() if k.endswith(".md")
    )
    for case, item in zip(native.cases, truth, strict=True):
        names = {upload_name(f["file"]) for f in item["evidence"]}
        assert names == {d.canonical_name for d in case.gold_documents}
        if item["category"] == "cross_document":
            assert len(names) == 2 and len(item["fact_ids"]) == 2
        if item["category"] == "no_answer":
            assert (
                not item["evidence"] and item["unknown_product_id"] not in corpus_text
            )
        if item["category"] == "version_policy":
            assert case.conflicts and len(item["evidence"]) >= 2


def build_assets(seed: int = SEED) -> dict[str, bytes]:
    records = make_records(seed)
    facts = make_facts(records)
    dataset, truth = make_cases(records, facts)
    corpus = {
        file: xlsx_bytes(*TABLES[file], rows)
        for file, rows in table_rows(records).items()
    }
    for file in MARKDOWN:
        lines = [
            f"# {Path(file).stem}",
            "",
            NOTICE,
            "",
            f"来源标识：nebula:{Path(file).stem}",
            "",
        ]
        for fact in facts:
            if fact["file"] == file:
                lines.extend([f"## {fact['section']}", "", fact["statement"], ""])
        corpus[file] = ("\n".join(lines) + "\n").encode("utf-8")
    validate_assets(records, dataset, truth, corpus)
    assets = {f"corpus/{name}": value for name, value in corpus.items()}
    for name, value in corpus.items():
        if name.endswith(".xlsx"):
            assets[f"upload_corpus/{name}"] = value
        else:
            lines = value.decode("utf-8").splitlines()
            html = (
                '<!doctype html><html><head><meta charset="utf-8"><title>'
                + escape(Path(name).stem)
                + "</title></head><body><main>\n"
            )
            for line in lines:
                if line:
                    tag = (
                        "h2"
                        if line.startswith("## ")
                        else "h1"
                        if line.startswith("# ")
                        else "p"
                    )
                    html += f"<{tag}>{escape(line.lstrip('# '))}</{tag}>\n"
            assets[f"upload_corpus/{upload_name(name)}"] = (
                html + "</main></body></html>\n"
            ).encode("utf-8")
    assets["source_records.json"] = json_bytes(records)
    assets["dataset_v1.json"] = json_bytes(dataset)
    for split in ("dev", "frozen"):
        subset = {
            **dataset,
            "name": f"nebula_{split}_v1",
            "cases": [c for c in dataset["cases"] if split in c["tags"]],
        }
        assets[f"dataset_{split}_v1.json"] = json_bytes(subset)
    assets["ground_truth.jsonl"] = (
        "\n".join(json.dumps(t, ensure_ascii=False, sort_keys=True) for t in truth)
        + "\n"
    ).encode("utf-8")
    assets["splits.json"] = json_bytes(
        {
            split: [t["case_id"] for t in truth if t["split"] == split]
            for split in ("dev", "frozen")
        }
    )
    sample_ids = [
        f"{kind}_{n:03d}"
        for kind in (
            "simple_fact",
            "table_query",
            "cross_document",
            "no_answer",
            "version_policy",
        )
        for n in (1, 10)
    ]
    review = [
        "# 人工复核清单",
        "",
        "以下为待人工勾选的 10 条代表 Case；自动一致性检查不替代人工复核。",
        "",
    ]
    for case_id in sample_ids:
        case = next(c for c in dataset["cases"] if c["id"] == case_id)
        item = next(t for t in truth if t["case_id"] == case_id)
        review.extend(
            [
                f"- [ ] {case_id}：{case['question']}",
                f"  参考：{case['reference_answer']}",
                "  来源："
                + "; ".join(
                    f["source_anchor"]
                    for f in item["evidence"]
                    or [{"source_anchor": "全语料不存在此型号；需检查拒答"}]
                ),
                "",
            ]
        )
    assets["manual_review.md"] = "\n".join(review).encode("utf-8")
    manifest = {
        "generator_version": VERSION,
        "seed": seed,
        "product_count": 200,
        "repair_count": 100,
        "document_count": 12,
        "case_count": 80,
        "split_counts": {"dev": 20, "frozen": 60},
        "dataset_fingerprint": dataset_fingerprint(
            RagEvalDataset.model_validate(dataset)
        ),
        "ingestion_allowlist": sorted(
            f"upload_corpus/{upload_name(n)}" for n in corpus
        ),
        "documents": [
            {
                "source_id": f"nebula:{Path(n).stem}",
                "source_file": f"corpus/{n}",
                "canonical_name": upload_name(n),
                "upload_file": f"upload_corpus/{upload_name(n)}",
            }
            for n in sorted(corpus)
        ],
        "files": {name: sha256(value) for name, value in sorted(assets.items())},
        "license": "Synthetic fixture data: CC0-1.0; generator follows repository code license.",
    }
    assets["manifest.json"] = json_bytes(manifest)
    return assets


def generate(output: Path, seed: int = SEED, *, check: bool = False) -> dict:
    assets = build_assets(seed)
    # Refuse differing files, including hand-edited generated assets. Use a new directory for new seeds.
    for name, value in assets.items():
        path = output / name
        if path.is_symlink() or any(p.is_symlink() for p in path.parents):
            raise ValueError(f"symlink output is not supported: {name}")
        if path.exists() and path.read_bytes() != value:
            raise ValueError(f"refusing to overwrite differing data: {name}")
        if check and not path.is_file():
            raise ValueError(f"missing generated asset: {name}")
    for folder in ("corpus", "upload_corpus"):
        directory = output / folder
        if directory.exists():
            actual = {
                p.relative_to(output).as_posix()
                for p in directory.rglob("*")
                if p.is_file()
            }
            assert actual <= set(assets), "unexpected file in ingestion corpus"
    if not check:
        for name, value in assets.items():
            path = output / name
            if not path.exists():
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(value)
    return json.loads(assets["manifest.json"])


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--seed", type=int, default=SEED)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    try:
        manifest = generate(args.output, args.seed, check=args.check)
    except (AssertionError, ValueError, OSError) as exc:
        print(f"generation failed: {exc}", file=sys.stderr)
        return 1
    print(
        json.dumps(
            {
                k: manifest[k]
                for k in (
                    "document_count",
                    "case_count",
                    "split_counts",
                    "dataset_fingerprint",
                )
            },
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
